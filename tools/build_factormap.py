# -*- coding: utf-8 -*-
"""FactorMap 데이터 v3 — Ver1 엑셀 기준
   · t=0 평가사례(선별분·지대 포함) / t=1 거래사례(나지) / t=2 거래사례(배분, 판정 O·△)
   · 개별 필지 좌표(fm_geo_cache) + 지대 + 주소
   · 기존 프론트 스키마(t·q·metric) 유지, j(지대)·a(주소)·ll(좌표) 추가"""
import sys, os, re, json, time
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
from collections import Counter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from geocode_fm import clean_jibun

SRC = r'\\server_new\공용폴더\2. 업무Part\♣해남 솔라시도\원준\가격자료(평가사례, 거래사례, 표준지)\거래사례, 평가사례_Ver1.xlsx'
OUT = r'D:\Dropbox\■ 전산TF\Github Clone\FactorMap_260815\data'
HERE = os.path.dirname(os.path.abspath(__file__))
GEO = json.load(open(os.path.join(HERE, 'fm_geo_cache.json'), encoding='utf-8'))
STD_SRC = ('//server_new/공용폴더/2. 업무Part/♣해남 솔라시도/원준/★솔라시도_GIS공간분석_2261필지_v5(260820).xlsx')
STD_GEO = json.load(open(os.path.join(HERE, 'std_geo.json'), encoding='utf-8'))

ZONES = ['일반상업지역', '제1종전용주거지역', '준공업지역', '제3종일반주거지역',
         '제1종일반주거지역', '일반공업지역', '준주거지역', '제2종전용주거지역',
         '제2종일반주거지역', '생산녹지지역', '계획관리지역', '자연녹지지역',
         '농림지역', '보전관리지역', '생산관리지역', '보전녹지지역', '근린상업지역',
         '개발제한구역', '자연환경보전지역', '기타·미상']
# 지목 그룹 — 사용자 엑셀 수식 기준 (정확일치 · TRIM)
#   제외 그룹(도로·구거·묘지 등)은 수록하지 않는다
USES = ['대지', '전답', '임야', '확인필요']
JD = ['대지', '전답', '임야', '확인필요']
G_DAEJI = {'공','공장','공장용','공장용지','대','대지','잡','잡종지','장','종','종교용지',
           '주','주유소','주유소 용지','주유소용','주유소용지','주차장','차','창','창고',
           '창고용지','체','체육용지','학'}
G_JEONDAP = {'과','과수원','답','목','목장','목장용지','유','유원지','전'}
G_IMYA = {'임','임야','임야외'}
G_JEOE = {'구','구거','도','도로','묘','묘지','수도용지','양','양어장','염','염전',
          '유지','제','제방','하천'}

ZMAP = {}
for z in ZONES[:-1]:
    ZMAP[z] = z
    ZMAP[z.replace('지역', '').replace('구역', '')] = z
ZMAP.update({
    '일상': '일반상업지역', '1전주': '제1종전용주거지역', '2전주': '제2종전용주거지역',
    '1주': '제1종일반주거지역', '1종일주': '제1종일반주거지역', '제1종일반': '제1종일반주거지역',
    '2주': '제2종일반주거지역', '2종일주': '제2종일반주거지역', '제2종일반': '제2종일반주거지역',
    '3주': '제3종일반주거지역', '3종일주': '제3종일반주거지역', '제3종일반': '제3종일반주거지역',
    '일공': '일반공업지역', '준공': '준공업지역',
    '자연': '자연녹지지역', '생산': '생산녹지지역',
    '계관': '계획관리지역', '생관': '생산관리지역', '보관': '보전관리지역',
    '농림': '농림지역', '근상': '근린상업지역', '개발제한': '개발제한구역',
    '자연환경': '자연환경보전지역',
    '1종전주': '제1종전용주거지역', '2종전주': '제2종전용주거지역',
})



def fnum(v):
    try:
        n = float(str(v).replace(',', ''))
        return n if n > 0 else None
    except Exception:
        return None


def zidx(s):
    s = re.sub(r'\s+', '', str(s or ''))
    z = ZMAP.get(s) or ZMAP.get(s.replace('지역', ''))
    return ZONES.index(z) if z else len(ZONES) - 1


def uidx(v):
    """지목 → 그룹 인덱스. 제외 그룹은 -1 (수록 안 함)"""
    t = str(v or '').strip()
    if t in G_DAEJI:
        return 0
    if t in G_JEONDAP:
        return 1
    if t in G_IMYA:
        return 2
    if t in G_JEOE:
        return -1
    return 3


def jidx(s):
    return JD.index(s) if s in JD else 3


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    rows, drop = [], Counter()
    purposes, pidx = [], {}

    def puidx(v):
        t = re.sub(r'\s+', ' ', str(v or '')).strip()
        if not t:
            return None
        if t not in pidx:
            pidx[t] = len(purposes)
            purposes.append(t)
        return pidx[t]

    # ── t=0 평가사례 (Ver1 선별분) ──
    ws = wb['평가사례']
    for r in range(2, ws.max_row + 1):
        so = re.sub(r'\s+', ' ', str(ws.cell(r, 4).value or '')).strip()
        toks = so.split()
        if len(toks) < 4:
            drop['평가:소재지 형식'] += 1
            continue
        jd = str(ws.cell(r, 10).value or '').strip()
        if jd in ('제외', ''):
            drop['평가:지대 제외/공란'] += 1
            continue
        rt = fnum(ws.cell(r, 17).value)
        up = fnum(ws.cell(r, 6).value)
        gp = fnum(ws.cell(r, 16).value)
        if not rt and up and gp:
            rt = up / gp
        if not rt and not up:
            drop['평가:단가·비율 없음'] += 1
            continue
        jb = clean_jibun(ws.cell(r, 5).value)
        y = str(ws.cell(r, 2).value or '')[:4]
        u = uidx(ws.cell(r, 18).value)
        if u < 0:
            drop['평가:지목 제외그룹(도로 등)'] += 1
            continue
        rows.append({'t': 0, 'pu': puidx(ws.cell(r, 1).value),
                     's': toks[1], 'e': toks[2], 'l': toks[3],
                     'z': zidx(ws.cell(r, 20).value), 'u': u,
                     'j': jidx(jd), 'y': int(y) if y.isdigit() else 0,
                     'r': round(rt, 3) if rt else None,
                     'p': int(up) if up else None, 'g': int(gp) if gp else None,
                     'q': 0, 'a': so + (' ' + jb if jb else ''),
                     'll': GEO.get(so + ' ' + jb) if jb else None})

    # ── 거래사례 시트 인덱스 (배분행 속성 참조용) + t=1 나지 ──
    ws = wb['거래사례']
    # 건물행(V 주용도 · W 주구조 · AA 단가표분류) — 배분 팝업에 건물 정보를 붙이기 위함
    bldg_attr = {}
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 9).value or '').strip() != '건물':
            continue
        bldg_attr[(str(ws.cell(r, 1).value or '').strip(),
                   str(ws.cell(r, 4).value or '').strip())] = {
            'use': str(ws.cell(r, 22).value or '').strip(),
            'st': str(ws.cell(r, 23).value or '').strip(),
            'cls': str(ws.cell(r, 27).value or '').strip(),
            'cst': str(ws.cell(r, 28).value or '').strip(),
        }
    land_attr = {}
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 9).value or '').strip() != '토지':
            continue
        sgg = str(ws.cell(r, 2).value or '').strip()
        loc = str(ws.cell(r, 3).value or '').strip()
        if not sgg or len(loc.split()) != 2:
            drop['거래:소재지 형식'] += 1
            continue
        no = str(ws.cell(r, 1).value or '').strip()
        jbraw = str(ws.cell(r, 4).value or '').strip()
        jd = str(ws.cell(r, 16).value or '').strip()
        attr = {'sgg': sgg, 'loc': loc, 'z': zidx(ws.cell(r, 8).value),
                'u': uidx(ws.cell(r, 7).value), 'jd': jd,
                'g': fnum(ws.cell(r, 18).value),
                'y': str(ws.cell(r, 10).value or '')[:4]}
        land_attr[(no, jbraw)] = attr

        if jd == '제외':
            drop['거래:지대 제외'] += 1
            continue
        up = fnum(ws.cell(r, 14).value)
        if not up:
            drop['거래:단가 없음'] += 1
            continue
        if attr['u'] < 0:
            drop['거래:지목 제외그룹(도로 등)'] += 1
            continue
        if '배분' in str(ws.cell(r, 36).value or ''):
            continue                          # 배분 연동분은 t=2 로 (배분법 시트에서)
        jb = clean_jibun(jbraw)
        emd, ri = loc.split()
        y = attr['y']
        rows.append({'t': 1, 's': sgg.split()[-1], 'e': emd, 'l': ri,
                     'z': attr['z'], 'u': attr['u'], 'j': jidx(jd),
                     'y': int(y) if y.isdigit() else 0,
                     'r': round(up / attr['g'], 3) if attr['g'] else None,
                     'p': int(up), 'g': int(attr['g']) if attr['g'] else None,
                     'q': 0, 'a': sgg + ' ' + loc + (' ' + jb if jb else ''),
                     'll': GEO.get(sgg + ' ' + loc + ' ' + jb) if jb else None})

    # ── t=2 배분 (배분법 시트 · 판정 O/△) ──
    al = wb['배분법']
    HR = 7
    for r in range(HR + 1, al.max_row + 1):
        no = str(al.cell(r, 1).value or '').strip()
        if not no:
            continue
        verdict = str(al.cell(r, 20).value or '').strip()
        if verdict.startswith('O'):
            q = 0
        elif verdict.startswith('△'):
            q = 1
        else:
            drop['배분:판정 X·계산불가'] += 1
            continue
        up = fnum(al.cell(r, 15).value)
        if not up:
            drop['배분:단가 없음'] += 1
            continue
        sgg = str(al.cell(r, 2).value or '').strip()
        loc = str(al.cell(r, 3).value or '').strip()
        jbraw = str(al.cell(r, 4).value or '').strip()
        attr = land_attr.get((no, jbraw)) or {}
        if attr.get('jd') == '제외':
            drop['배분:지대 제외'] += 1
            continue
        if len(loc.split()) != 2:
            drop['배분:소재지 형식'] += 1
            continue
        emd, ri = loc.split()
        jb = clean_jibun(jbraw)
        rt = fnum(al.cell(r, 16).value)
        y = str(al.cell(r, 5).value or '')[:4]
        if attr.get('u', 3) < 0:
            drop['배분:지목 제외그룹'] += 1
            continue
        # 배분 근거 — 팝업에서 '무슨 값으로 건물을 빼고 토지단가를 냈는지' 보여준다
        d = {
            'amt': fnum(al.cell(r, 6).value),      # 거래금액
            'la': fnum(al.cell(r, 7).value),       # 토지 거래면적
            'ba': fnum(al.cell(r, 8).value),       # 건물 공부면적
            'uc': fnum(al.cell(r, 9).value),       # 재조달 표준단가
            'life': fnum(al.cell(r, 10).value),    # 내용연수
            'age': fnum(al.cell(r, 11).value),     # 경과연수
            'res': fnum(al.cell(r, 12).value),     # 잔가율
            'bp': fnum(al.cell(r, 13).value),      # 건물가격
            'lp': fnum(al.cell(r, 14).value),      # 토지배분액
            'gr': str(al.cell(r, 18).value or '').strip(),   # 실제 적용급수
        }
        d = {k: (round(v, 4) if isinstance(v, float) else v) for k, v in d.items() if v}
        b = bldg_attr.get((no, jbraw))
        if b:
            d.update({k: v for k, v in b.items() if v})   # use/st/cls/cst
        rows.append({'t': 2, 's': sgg.split()[-1], 'e': emd, 'l': ri,
                     'z': attr.get('z', len(ZONES) - 1), 'u': attr.get('u', 3),
                     'j': jidx(attr.get('jd', '')),
                     'y': int(y) if y.isdigit() else 0,
                     'r': round(rt, 3) if rt else None,
                     'd': d,
                     'p': int(up), 'g': int(attr['g']) if attr.get('g') else None,
                     'q': q, 'a': sgg + ' ' + loc + (' ' + jb if jb else ''),
                     'll': GEO.get(sgg + ' ' + loc + ' ' + jb) if jb else None})

    # ── t=3 표준지 ((Raw)본건외 표준지 · 지대 '제외' 뺌) ──
    ws = openpyxl.load_workbook(STD_SRC, data_only=True, read_only=True)['(Raw)본건외 표준지']
    JD_STD = {'대지': 0, '전답': 1, '임야': 2}
    for row in ws.iter_rows(min_row=2, values_only=True):
        so = re.sub(r'\s+', ' ', str(row[1] or '')).strip()
        jb = str(row[2] or '').strip()
        if not so or not jb:
            continue
        jd = str(row[4] or '').strip()
        if jd == '제외' or jd not in JD_STD:
            drop['표준지:지대 제외/불명'] += 1
            continue
        addr = so + ' ' + jb
        geo = STD_GEO.get(addr) or {}
        toks = so.split()
        if len(toks) == 1:                 # 목포 시내 동(洞) — 읍면·리 자리에 동명
            toks = [so, so]
        elif len(toks) != 2:
            drop['표준지:소재지 형식'] += 1
            continue
        y = str(row[0] or '')[:4]
        gp = fnum(row[9])
        sp = {}
        for k2, idx2 in [('jm', 3), ('us', 6), ('z2', 8), ('rd', 10), ('sh', 11),
                         ('pf', 14), ('pr', 15), ('et', 17)]:
            v = str(row[idx2] or '').strip()
            if v and v.lower() not in ('none',):
                sp[k2] = v
        rows.append({'t': 3, 's': geo.get('sgg') or '미상', 'e': toks[0], 'l': toks[1],
                     'z': zidx(row[7]), 'u': JD_STD[jd], 'j': JD_STD[jd],
                     'y': int(y) if y.isdigit() else 0,
                     'r': None, 'p': None, 'g': int(gp) if gp else None,
                     'q': 0, 'a': addr, 'll': geo.get('ll'), 'sp': sp})

    n_by_t = Counter(x['t'] for x in rows)
    print(f'수록 {len(rows):,} (평가 {n_by_t[0]:,} · 거래나지 {n_by_t[1]:,} · 배분 {n_by_t[2]:,}'
          f' · 표준지 {n_by_t[3]:,})')
    print(f'좌표 있음 {sum(1 for x in rows if x["ll"]):,} · 개공비율 있음 {sum(1 for x in rows if x["r"]):,}')
    print('제외:', dict(drop.most_common()))

    # 리 중심 — 기존 캐시 재사용, 신규 리는 소속 필지 평균
    old = {}
    p_old = os.path.join(OUT, 'records.js')
    if os.path.exists(p_old):
        m = re.search(r'window\.FM_CENTERS = (.*?);\n', open(p_old, encoding='utf-8').read())
        if m:
            old = json.loads(m.group(1))
    centers, acc = {}, {}
    for x in rows:
        k = x['s'] + '|' + x['e'] + '|' + x['l']
        if k in old:
            centers[k] = old[k]
        elif x['ll']:
            a = acc.setdefault(k, [0.0, 0.0, 0])
            a[0] += x['ll'][0]
            a[1] += x['ll'][1]
            a[2] += 1
    for k, a in acc.items():
        if k not in centers:
            centers[k] = [round(a[0] / a[2], 6), round(a[1] / a[2], 6)]
    print(f'리 {len(centers)}곳')

    meta = {'zones': ZONES, 'nzchip': 12, 'uses': USES, 'nuchip': 4, 'jd': JD,
            'types': ['평가사례', '거래사례', '거래사례(배분)', '표준지'],
            'purposes': purposes,
            'alloc': n_by_t[2], 'built': time.strftime('%Y-%m-%d %H:%M'), 'n': len(rows)}
    with open(os.path.join(OUT, 'records.js'), 'w', encoding='utf-8') as f:
        f.write('window.FM_META = ' + json.dumps(meta, ensure_ascii=False) + ';\n')
        f.write('window.FM_ROWS = ' + json.dumps(rows, ensure_ascii=False, separators=(',', ':')) + ';\n')
        f.write('window.FM_CENTERS = ' + json.dumps(centers, ensure_ascii=False) + ';\n')
    print(f'저장: records.js ({os.path.getsize(os.path.join(OUT, "records.js"))//1024} KB)')

    # 데이터 갱신이 브라우저 캐시에 막히지 않도록 script 태그 버전 쿼리를 갱신한다
    # (GitHub Pages 는 subresource 를 오래 캐시해서, 새로고침해도 옛 데이터가 붙는다)
    idx = os.path.join(os.path.dirname(OUT), 'index.html')
    if os.path.exists(idx):
        h = open(idx, encoding='utf-8').read()
        stamp = time.strftime('%Y%m%d%H%M')
        h2 = re.sub(r'(<script src="data/(?:records|boundaries|parcels|cart_shared)\.js)(\?v=\d+)?(">)',
                    lambda m: m.group(1) + '?v=' + stamp + m.group(3), h)
        if h2 != h:
            open(idx, 'w', encoding='utf-8').write(h2)
            print(f'index.html 데이터 버전 갱신: ?v={stamp}')


if __name__ == '__main__':
    main()
