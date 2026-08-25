# -*- coding: utf-8 -*-
"""records.js 에 시점(연월일) 필드 dt 주입
   전체 재빌드(build_factormap.py)에는 지오코딩 캐시가 필요해 회사 PC에서만 가능하다.
   이 스크립트는 기존 records.js 를 그대로 두고 dt 만 채워 넣는다.
   다음 정식 재빌드 시 build_factormap.py 가 같은 값을 생성하므로 결과는 동일하다."""
import sys, os, re, json, datetime
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from collections import Counter, defaultdict

SRC = ('//server_new/공용폴더/2. 업무Part/♣해남 솔라시도/원준/'
       '가격자료(평가사례, 거래사례, 표준지)/거래사례, 평가사례_Ver1.xlsx')
REC = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'records.js')


def fdate(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    t = str(v or '').strip()
    m = re.match(r'(\d{4})[-./년]\s*(\d{1,2})[-./월]\s*(\d{1,2})', t)
    return '%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3))) if m else None


def fnum(v):
    try:
        n = float(str(v).replace(',', ''))
        return n if n > 0 else None
    except Exception:
        return None


wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

# ── 평가사례 : (연도, 단가, 공시지가) → 기준시점 ──
ev = defaultdict(set)
ws = wb['평가사례']
for r in ws.iter_rows(min_row=2, values_only=True):
    dt = fdate(r[11])                       # L 기준시점
    if not dt:
        continue
    y = str(r[1] or '')[:4]
    up, gp = fnum(r[5]), fnum(r[15])
    if y.isdigit():
        ev[(int(y), int(up) if up else None, int(gp) if gp else None)].add(dt)

# ── 거래사례 : (연도, 거래단가, 공시지가) → 거래시점 ──
tr = defaultdict(set)
ws = wb['거래사례']
for r in ws.iter_rows(min_row=2, values_only=True):
    dt = fdate(r[9])                        # J 거래시점
    if not dt:
        continue
    y = dt[:4]
    up, gp = fnum(r[13]), fnum(r[17])
    tr[(int(y), int(up) if up else None, int(gp) if gp else None)].add(dt)

print('평가 키 %d (충돌 %d) · 거래 키 %d (충돌 %d)'
      % (len(ev), sum(1 for v in ev.values() if len(v) > 1),
         len(tr), sum(1 for v in tr.values() if len(v) > 1)))

# ── records.js ──
raw = open(REC, encoding='utf-8').read()
m = re.search(r'window\.FM_ROWS = (\[.*?\]);\n', raw, re.S)
rows = json.loads(m.group(1))
print('records.js 행 %d' % len(rows))

hit = Counter()
amb = 0
for x in rows:
    y, p, g = x.get('y'), x.get('p'), x.get('g')
    if x['t'] == 3:                                     # 표준지 = 공시기준일 1/1
        x['dt'] = ('%d-01-01' % y) if y else None
        hit['표준지'] += 1 if y else 0
        continue
    src = ev if x['t'] == 0 else tr
    cand = src.get((y, p, g))
    if cand and len(cand) == 1:
        x['dt'] = next(iter(cand))
        hit['t%d' % x['t']] += 1
    elif cand:
        x['dt'] = sorted(cand)[0]                       # 동일 (연도·단가·공시가) 복수 → 가장 이른 날
        hit['t%d' % x['t']] += 1
        amb += 1
    else:
        x['dt'] = None

tot = sum(hit.values())
print('dt 채움 %d / %d (%.1f%%) — 유형별 %s · 다중후보 %d'
      % (tot, len(rows), tot * 100 / len(rows), dict(hit), amb))
miss = Counter(x['t'] for x in rows if not x.get('dt'))
print('미매칭:', dict(miss))

out = raw[:m.start(1)] + json.dumps(rows, ensure_ascii=False, separators=(',', ':')) + raw[m.end(1):]
open(REC, 'w', encoding='utf-8').write(out)
print('저장: records.js (%d KB)' % (os.path.getsize(REC) // 1024))

# 캐시 무효화용 버전 스탬프 갱신
idx = os.path.join(os.path.dirname(REC), '..', 'index.html')
h = open(idx, encoding='utf-8').read()
stamp = datetime.datetime.now().strftime('%Y%m%d%H%M')
h2 = re.sub(r'(<script src="data/(?:records|boundaries|parcels|cart_shared)\.js)(\?v=\d+)?(">)',
            lambda mm: mm.group(1) + '?v=' + stamp + mm.group(3), h)
if h2 != h:
    open(idx, 'w', encoding='utf-8').write(h2)
    print('index.html 데이터 버전 갱신: ?v=%s' % stamp)
