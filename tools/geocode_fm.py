# -*- coding: utf-8 -*-
"""FactorMap 필지 지오코딩 — geocoder.py(Skills) 방식: 동시 8워커·QPS 9·캐시 이어하기

   · 평가사례(D 소재지 + E 지번) · 거래사례 토지행(B 시군구 + C 리 + D 지번) 의
     주소를 모아 카카오 주소검색으로 좌표를 얻어 fm_geo_cache.json 에 쌓는다.
   · build_factormap.py 가 이 캐시(GEO)와 clean_jibun() 을 그대로 쓴다.
     → clean_jibun 의 결과는 사례 주소 문자열(records.js 의 a)이 되고, 그 주소가
       장바구니 키의 일부이므로 규칙을 바꾸면 저장된 장바구니가 깨진다. 주의.
   · 캐시 값: 주소 → [위도, 경도] · 카카오가 못 찾으면 None (재조회 안 함)

   [260826] 원본 .py 유실 → 컴파일본(.pyc) 역어셈블로 복원. 동작 동일함을 확인.
            SRC 는 원본 폴더 개명(… 거래사례, 표준지)에 맞춰 갱신.
"""
import sys
import os
import re
import json
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

SRC = (r'\\server_new\공용폴더\2. 업무Part\♣해남 솔라시도\원준'
       r'\가격자료(평가사례, 거래사례, 표준지)\거래사례, 평가사례_Ver1.xlsx')
HERE = os.path.dirname(os.path.abspath(__file__))
CACHE_FILE = os.path.join(HERE, 'fm_geo_cache.json')

KEY = '8b9a82bbb05d02a6b63170a7c257bf7d'
URL = 'https://dapi.kakao.com/v2/local/search/address.json'
QPS = 9.0
MAX_WORKERS = 8
SAVE_EVERY = 500


def clean_jibun(v):
    """'  1705-   1' → '1705-1', '산  12- 3' → '산12-3'. 유효 지번 아니면 ''"""
    s = re.sub(r'\s+', '', str(v or ''))
    s = s.replace('번지', '')
    m = re.match(r'^(산)?(\d+)(?:-(\d+))?$', s)
    if not m or int(m.group(2)) == 0:
        return ''
    return (m.group(1) or '') + m.group(2) + \
           ('-' + str(int(m.group(3))) if m.group(3) and int(m.group(3)) else '')


def collect():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    addrs = set()
    ws = wb['평가사례']
    for r in range(2, ws.max_row + 1):
        so = re.sub(r'\s+', ' ', str(ws.cell(r, 4).value or '')).strip()
        jb = clean_jibun(ws.cell(r, 5).value)
        if not so or not jb:
            continue
        addrs.add(so + ' ' + jb)
    ws = wb['거래사례']
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 9).value or '').strip() != '토지':
            continue
        sgg = str(ws.cell(r, 2).value or '').strip()
        loc = str(ws.cell(r, 3).value or '').strip()
        jb = clean_jibun(ws.cell(r, 4).value)
        if not sgg or not loc or not jb:
            continue
        addrs.add(sgg + ' ' + loc + ' ' + jb)
    return sorted(addrs)


class Limiter:
    """초당 rate 건으로 제한 (토큰 버킷)"""

    def __init__(self, rate):
        self.rate = self.tokens = rate
        self.last = time.monotonic()
        self.lock = threading.Lock()

    def acquire(self):
        with self.lock:
            now = time.monotonic()
            self.tokens = min(self.rate, self.tokens + (now - self.last) * self.rate)
            self.last = now
            if self.tokens >= 1:
                self.tokens -= 1
                wait = 0.0
            else:
                wait = (1 - self.tokens) / self.rate
                self.tokens = 0
                self.last = now + wait
        if wait > 0:
            time.sleep(wait)


def main():
    addrs = collect()
    cache = json.load(open(CACHE_FILE, encoding='utf-8')) if os.path.exists(CACHE_FILE) else {}
    todo = [a for a in addrs if a not in cache]
    print(f'유니크 주소 {len(addrs):,} · 미처리 {len(todo):,}')
    if not todo:
        print('전부 캐시에 있음')
    else:
        ses = requests.Session()
        lim = Limiter(QPS)
        lock = threading.Lock()
        done = [0, 0, 0]

        def one(a):
            for att in range(3):
                lim.acquire()
                try:
                    r = ses.get(URL,
                                headers={'Authorization': 'KakaoAK ' + KEY},
                                params={'query': a, 'analyze_type': 'similar'},
                                timeout=10)
                    r.raise_for_status()
                    d = r.json().get('documents') or []
                    return a, ([round(float(d[0]['y']), 6), round(float(d[0]['x']), 6)]
                               if d else None), True
                except Exception:
                    time.sleep(1.5 * (att + 1))
            return a, None, False

        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
            for fut in as_completed({ex.submit(one, a) for a in todo}):
                a, v, ok = fut.result()
                with lock:
                    if ok:
                        cache[a] = v
                    done[0] += 1
                    done[1] += 1 if (ok and v) else 0
                    done[2] += 0 if ok else 1
                    if done[0] % SAVE_EVERY == 0:
                        json.dump(cache, open(CACHE_FILE, 'w', encoding='utf-8'),
                                  ensure_ascii=False)
                        print(f'  {done[0]:,}/{len(todo):,}'
                              f' (성공 {done[1]:,} · 네트오류 {done[2]:,})')
        json.dump(cache, open(CACHE_FILE, 'w', encoding='utf-8'), ensure_ascii=False)

    hit = sum(1 for a in addrs if cache.get(a))
    miss = sum(1 for a in addrs if a in cache and cache[a] is None)
    print(f'완료 — 좌표 확보 {hit:,} · 카카오 미스 {miss:,} · 미처리 {len(addrs) - hit - miss:,}')


if __name__ == '__main__':
    main()
